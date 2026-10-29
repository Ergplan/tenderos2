"""Build docs/enterprise/FEATURES-JOULEWISE-CLIENT.xlsx: the enterprise feature list with, for every
feature, what jouleWise does (builds, ships, patches under AMC, supports) and what the client does
(provides, enters, reviews, operates). Reads the list in build_enterprise_features.py. Run from
the repository root:
    python3 tools/build_features_split.py
"""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parent))
from build_enterprise_features import PRIORITY, STATUS, TEAMS, ZONES, rows  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "enterprise" / "FEATURES-JOULEWISE-CLIENT.xlsx"

# What each side does, by area; a feature can override either side below
AREA_SPLIT = {
    "Tender intake": (
        "Builds and maintains the reading engine (parsers, extraction models, prompts) and the intake-zone screens; ships fixes under AMC",
        "Uploads the tender documents in the intake zone; Legal reviews the reading (every value on its page) before it is packaged for the core",
    ),
    "Bid workspace": (
        "Builds, tests and patches the bid workspace; keeps the steps, benchmarks and tender packs current",
        "Projects, Design, Finance and Legal enter their inputs step by step, attach evidence and sign off; the tender manager runs the bid",
    ),
    "Design": (
        "Builds the design tools (sources, studios, profiles, schematic, losses, DSM model) and keeps the engineering norms current",
        "Design team provides the CUF files, costs, site layout and losses, with evidence; validates the DSM norms against the regulations in force",
    ),
    "Sizing": (
        "Owns the optimizer (jouleWise IP, delivered sealed), its performance and its tuning; supports runs under AMC",
        "Runs the sizing on its own hardware; reviews the sizes and the compliance with the tender",
    ),
    "Finance": (
        "Builds the 25-year financial model and its Excel twin, checked against LibreOffice; keeps tax and CERC norms current",
        "Finance team provides the financing terms, lender's conditions and working-capital norms, with evidence; validates against its own model",
    ),
    "What-if": (
        "Builds the What-if and the tariff levers; keeps them in step with the model",
        "Management and Finance use them on bid day to decide the tariff to quote",
    ),
    "Evidence": (
        "Builds evidence, sign-off, the register and the audit log; patches them under AMC",
        "Every team attaches evidence to its inputs; the super admin signs values off; IT keeps the exports",
    ),
    "Platform": (
        "Builds and signs the installation bundles and patches; supports installation and upgrades under AMC",
        "IT installs on its own hardware, applies patches, runs backups, manages users and the firewall",
    ),
}
# per-feature overrides: feature name -> (jouleWise, client); None keeps the area's text
FEATURE_SPLIT = {
    "Claude connector (MCP)": ("Provides the connector and its tools; off in the core by default", "Decides whether the connector is allowed in the intake zone; manages the token"),
    "Transfer package to the core": ("Builds the signed package format, the export and the import tool", "IT carries the package across the one-way transfer (file drop or diode) and imports it in the core"),
    "Review before transfer": ("Builds the review screen with every value on its page", "Legal reviews and confirms the reading before transfer"),
    "Users, sign-in and tab access": ("Builds accounts, roles and tab access", "IT creates users, grants tabs and removes leavers"),
    "Enterprise sign-in (SSO)": ("Integrates SAML/OIDC or LDAP", "IT provides the identity provider, the metadata and the groups"),
    "Air-gapped installation": ("Builds the offline bundle (images, compose, migrations, checksums, signature) and the installer", "IT provides the servers, Docker, the certificate and the firewall rules; installs with install.sh"),
    "Patches under AMC": ("Signs and ships each patch bundle with its changelog; supports the apply and any rollback", "IT verifies and applies the bundle with tenderos-patch; keeps the backups it makes"),
    "Backups and restore": ("Provides the backup and restore scripts and tests them", "IT schedules the backups, stores them off the server and tests a restore"),
    "HTTPS and firewall": ("Configures TLS at the edge and documents the rules per zone", "IT provides the certificate and enforces the firewall allowlists"),
    "Day and night mode": ("Builds it", "Each user picks"),
    "Demo mode": ("Provides it for demonstrations", "Not used in the core: real sign-in only"),
    "Team setup and roles": ("Builds Team setup", "The super admin adds people, roles and the tabs they cover"),
    "Benchmarks": ("Keeps the CERC norms and market prices current in the benchmarks", "Replaces benchmarks with its own numbers and evidence"),
    "Bid forms and letters": ("Builds every format and letter of each tender pack", "Legal reviews, edits the wording and signs the pack"),
    "Deviation settlement (DSM) per CERC": ("Builds the DSM model with editable norms and keeps the regulations' structure current", "Design and Finance set the forecast-error and the norms from the notified text; validate against their DSM bills"),
    "25-year CUF justification": ("Builds the justification note and the Word export", "Design provides the data source, period, P50/P90 report and the assessor's name"),
    "CUF file download": ("Builds the download", "Design keeps the file with the bid pack as evidence"),
}

HEAD = ["ID", "Area", "Feature", "What it does", "Client team", "Zone", "Status", "Priority", "jouleWise", "Client", "Notes / where"]
WIDTHS = [8, 14, 30, 60, 12, 8, 9, 9, 48, 48, 40]


def split_rows():
    out = []
    for r in rows():
        rid, area, feat, what, team, zone, status, prio, notes = r
        jw, cl = AREA_SPLIT.get(area, ("Builds and maintains it", "Uses it"))
        o = FEATURE_SPLIT.get(feat)
        if o:
            jw = o[0] or jw
            cl = o[1] or cl
        out.append([rid, area, feat, what, team, zone, status, prio, jw, cl, notes])
    return out


def build(path: Path):
    wb = Workbook()
    font = Font(name="Arial", size=10)
    bold = Font(name="Arial", size=10, bold=True)
    title = Font(name="Arial", size=14, bold=True)
    head_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="1F3864")
    jw_fill = PatternFill("solid", fgColor="E2EFDA")  # jouleWise column
    cl_fill = PatternFill("solid", fgColor="FFF2CC")  # client column
    area_fill = PatternFill("solid", fgColor="D9E1F2")
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="top")
    center = Alignment(horizontal="center", vertical="top", wrap_text=True)

    # ---- README
    ws = wb.active
    ws.title = "README"
    lines = [
        ("TenderOS Enterprise: features by jouleWise and client responsibility", title),
        ("", font),
        ("One row per feature on the Features sheet. For each: what jouleWise does (green column) and what the client does (yellow column).", font),
        ("jouleWise: builds, tests, signs and ships the software, keeps the norms and tender packs current, and patches under the AMC.", font),
        ("Client: provides the inputs and evidence, runs the installation on its own hardware, reviews and signs off, and submits the bid.", font),
        ("", font),
        ("Columns", bold),
        ("ID · Area · Feature · What it does · Client team (which client team fills it in or owns it) · Zone · Status · Priority · jouleWise · Client · Notes / where", font),
        ("", font),
        ("Zone", bold),
    ] + [(f"{k}: {v}", font) for k, v in ZONES.items()] + [("", font), ("Status", bold)] + [(f"{k}: {v}", font) for k, v in STATUS.items()] + [("", font), ("Priority", bold)] + [(f"{k}: {v}", font) for k, v in PRIORITY.items()] + [("", font), ("Client team", bold)] + [(f"{k}: {v}", font) for k, v in TEAMS.items()] + [
        ("", font),
        ("The Summary sheet counts the Features sheet with formulas (by area and status, and by client team); it follows any edit.", font),
        ("Status and Priority cells on the Features sheet have drop-down lists. Edit the jouleWise and Client text freely.", font),
        ("Built by tools/build_features_split.py from the list in tools/build_enterprise_features.py.", font),
    ]
    for i, (text, f) in enumerate(lines, 1):
        c = ws.cell(row=i, column=1, value=text)
        c.font = f
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 130

    # ---- Features
    data = split_rows()
    ws = wb.create_sheet("Features")
    ws.append(HEAD)
    for c in ws[1]:
        c.font = head_font
        c.fill = head_fill
        c.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
        c.border = border
    ws.row_dimensions[1].height = 30
    for r in data:
        ws.append(r)
    n = ws.max_row
    for row in ws.iter_rows(min_row=2, max_row=n):
        for c in row:
            c.font = font
            c.alignment = wrap
            c.border = border
        row[1].fill = area_fill
        row[8].fill = jw_fill
        row[9].fill = cl_fill
        for k in (0, 4, 5, 6, 7):
            row[k].alignment = center
    for i, w in enumerate(WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEAD))}{n}"
    dv_status = DataValidation(type="list", formula1='"Live,Extend,Build"', allow_blank=False, showErrorMessage=True, errorTitle="Status", error="One of Live, Extend, Build")
    dv_prio = DataValidation(type="list", formula1='"P1,P2,P3"', allow_blank=False, showErrorMessage=True, errorTitle="Priority", error="One of P1, P2, P3")
    ws.add_data_validation(dv_status)
    ws.add_data_validation(dv_prio)
    dv_status.add(f"G2:G{n}")
    dv_prio.add(f"H2:H{n}")
    # status colours follow the cell's text
    ws.conditional_formatting.add(f"G2:G{n}", FormulaRule(formula=['$G2="Live"'], fill=PatternFill("solid", fgColor="C6EFCE"), font=Font(name="Arial", size=10, color="006100")))
    ws.conditional_formatting.add(f"G2:G{n}", FormulaRule(formula=['$G2="Extend"'], fill=PatternFill("solid", fgColor="FFEB9C"), font=Font(name="Arial", size=10, color="9C5700")))
    ws.conditional_formatting.add(f"G2:G{n}", FormulaRule(formula=['$G2="Build"'], fill=PatternFill("solid", fgColor="FFC7CE"), font=Font(name="Arial", size=10, color="9C0006")))
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"

    # ---- Summary (formulas on the Features sheet)
    ws = wb.create_sheet("Summary")
    ws["A1"] = "Features by area and status"
    ws["A1"].font = title
    ws.append([])
    ws.append(["Area", "Live", "Extend", "Build", "Total"])
    for c in ws[3]:
        c.font = head_font
        c.fill = head_fill
        c.border = border
    areas = list(dict.fromkeys(r[1] for r in data))
    r0 = 4
    for i, a in enumerate(areas):
        r = r0 + i
        ws.cell(row=r, column=1, value=a)
        for col, st in ((2, "Live"), (3, "Extend"), (4, "Build")):
            ws.cell(row=r, column=col, value=f'=COUNTIFS(Features!$B$2:$B${n},$A{r},Features!$G$2:$G${n},"{st}")')
        ws.cell(row=r, column=5, value=f"=SUM(B{r}:D{r})")
    rt = r0 + len(areas)
    ws.cell(row=rt, column=1, value="Total").font = bold
    for col in range(2, 6):
        L = get_column_letter(col)
        c = ws.cell(row=rt, column=col, value=f"=SUM({L}{r0}:{L}{rt - 1})")
        c.font = bold
    for row in ws.iter_rows(min_row=r0, max_row=rt, max_col=5):
        for c in row:
            c.border = border
            if c.column > 1:
                c.alignment = Alignment(horizontal="center")
            if c.font != bold:
                c.font = font

    r = rt + 2
    ws.cell(row=r, column=1, value="Features by client team").font = title
    r += 2
    ws.append([])
    for col, h in enumerate(["Client team", "Features", "Live", "To build or extend"], 1):
        c = ws.cell(row=r, column=col, value=h)
        c.font = head_font
        c.fill = head_fill
        c.border = border
    teams = list(dict.fromkeys(x[4] for x in data))
    for i, t in enumerate(teams):
        rr = r + 1 + i
        ws.cell(row=rr, column=1, value=t)
        ws.cell(row=rr, column=2, value=f"=COUNTIF(Features!$E$2:$E${n},$A{rr})")
        ws.cell(row=rr, column=3, value=f'=COUNTIFS(Features!$E$2:$E${n},$A{rr},Features!$G$2:$G${n},"Live")')
        ws.cell(row=rr, column=4, value=f"=B{rr}-C{rr}")
        for col in range(1, 5):
            c = ws.cell(row=rr, column=col)
            c.border = border
            c.font = font
            if col > 1:
                c.alignment = Alignment(horizontal="center")
    r_end = r + len(teams)
    ws.cell(row=r_end + 2, column=1, value="Responsibility, in one line").font = title
    ws.cell(row=r_end + 3, column=1, value="jouleWise").font = bold
    ws.cell(row=r_end + 3, column=1).fill = jw_fill
    ws.cell(row=r_end + 3, column=2, value="Builds, tests, signs and ships every feature; keeps norms, benchmarks and tender packs current; patches the core under the AMC; supports IT on installation, upgrades and rollback.").font = font
    ws.cell(row=r_end + 4, column=1, value="Client").font = bold
    ws.cell(row=r_end + 4, column=1).fill = cl_fill
    ws.cell(row=r_end + 4, column=2, value="Provides the tender documents, inputs, evidence and lender's terms; installs and runs TenderOS on its own hardware; reviews and signs off each step; decides and submits the bid.").font = font
    for rr in (r_end + 3, r_end + 4):
        ws.cell(row=rr, column=2).alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=rr, start_column=2, end_row=rr, end_column=5)
        ws.row_dimensions[rr].height = 42
    ws.column_dimensions["A"].width = 22
    for L in "BCDE":
        ws.column_dimensions[L].width = 16
    ws.freeze_panes = "A4"

    wb.save(path)
    return len(data)


if __name__ == "__main__":
    k = build(OUT)
    print(f"{k} features → {OUT}")
