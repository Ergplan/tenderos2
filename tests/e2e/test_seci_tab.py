"""Browser end to end: the SECI RTC tender tab.

* test_documents_and_conditions: the five documents (the RfS among them), how the bid goes, the
  amendment trail, SECI's clarifications, the points to raise with SECI and the bid conditions with
  their quotes and the RfS's original wording; a PDF link opens the page.
* test_bid_on_screen: the bid (200 MW, any 6 Peak Hours, pay the penalty, a PPA date and the
  dates it sets), the benchmarks for the sources, HiGHS sizing against every peak and off-peak
  block, the money at stake; a tariff quoted below the target IRR is warned with the equity and
  project IRR it earns; the particulars and the RfS formats are filled and reviewed one by one; the
  bid pack (Word) and the financial model (Excel, its Live model recalculated by LibreOffice when
  installed) download.
* test_agent_panel: the GPT-Live bid agent's panel; without a key on the server it says how to set
  it up; with one, starting it sends the browser's WebRTC offer and shows the voice service's
  answer (a refusal, against a stand-in for OpenAI).

Needs the whole stack running; see tests/e2e/test_tender_to_bid.py for the environment.
"""
from __future__ import annotations

import json
import re
import zipfile
from pathlib import Path

import pymupdf
import pytest

from tests.e2e.test_tender_to_bid import DEMO, EMAIL, PASSWORD, URL, _check_live_model, _session

ROOT = Path(__file__).resolve().parents[2]
PACK = json.loads((ROOT / "web" / "src" / "seci" / "seci-rtc-tm-v.json").read_text())

pytestmark = [
    pytest.mark.slow,
    pytest.mark.skipif(not (URL and (DEMO or (EMAIL and PASSWORD))), reason="set TENDEROS_URL and TENDEROS_EMAIL, TENDEROS_PASSWORD (or TENDEROS_DEMO=1)"),
]


def _open_tab(page):
    page.click("aside.sidebar button:has-text('SECI RTC tender')")
    page.wait_for_selector("[data-testid=seci-tab]", timeout=30_000)


def _review(page):
    """Mark the open format reviewed. The RfS page beside the blanks loads as the format opens and
    can move the button under the pointer, so click again until the review shows."""
    for _ in range(4):
        page.locator("[data-testid=seci-format-review]").scroll_into_view_if_needed()
        page.wait_for_timeout(300)
        page.click("[data-testid=seci-format-review]")
        try:
            page.wait_for_selector("[data-testid=seci-format-reviewed]", timeout=4_000)
            return
        except Exception:  # noqa: BLE001 - try again
            continue
    raise AssertionError("the review did not register")


def test_documents_and_conditions():
    with _session() as page:
        _open_tab(page)
        docs = page.locator("[data-testid=seci-docs] article")
        assert docs.count() == len(PACK["documents"]) == 5
        rfs = page.locator("[data-testid=seci-doc-rfs]")
        assert "139 pages" in rfs.inner_text() and rfs.locator("a:has-text('Open the PDF')").get_attribute("href").endswith("RfS.pdf")
        assert page.locator("[data-testid=seci-process] li").count() == len(PACK["bidProcess"]) >= 10
        assert page.locator("[data-testid=seci-queries] li").count() == len(PACK["queries"]) >= 10
        assert page.locator("[data-testid=seci-changes] .seci-change").count() == len(PACK["changes"])
        page.click("button:has-text('Amendment-02')")
        assert page.locator("[data-testid=seci-changes] .seci-change").count() == sum(1 for c in PACK["changes"] if c["doc"] == "amd2")
        dfr = page.locator("[data-testid=seci-change-a1-8\\.1a]")
        assert dfr.count() == 0, "filtered to Amendment-02"
        page.click("button:has-text('All')")
        text = page.locator("[data-testid=seci-change-a1-8\\.1a]").inner_text()
        assert "80%" in text and "70% in each time-block" in text, text
        assert page.locator("[data-testid=seci-clarifications] .seci-change").count() == len(PACK["clarifications"])
        # a quote's link opens its PDF at the page
        href = page.locator("[data-testid=seci-change-a2-16\\.1] a.seci-link").first.get_attribute("href")
        assert href.endswith("Amendment-02.pdf#page=3"), href
        r = page.request.get(URL + href.split("#")[0])  # the session's cookie: the tender PDFs are behind the login
        assert r.status == 200 and r.body()[:5] == b"%PDF-", r.status
        page.click("[data-testid=seci-step-conditions]")
        n = len(PACK["conditions"])
        assert page.locator(".seci-cond").count() == n >= 50
        summary = page.locator("[data-testid=seci-cond-summary]").inner_text()
        assert f"{n} conditions" in summary and "removed" in summary, summary
        offpeak = page.locator("[data-testid=seci-cond-dfrOffpeak]")
        assert "70%" in offpeak.inner_text()
        offpeak.locator("summary:has-text('RfS as issued')").click()
        assert "80%" in offpeak.inner_text(), "the RfS's original off-peak floor, before Amendment-01"
        assert "Unclear" in page.locator("[data-testid=seci-cond-technical]").inner_text()
        assert "25 years from the Scheduled Commencement" in page.locator("[data-testid=seci-cond-term]").inner_text()
        assert "None stated in the RfS" in page.locator("[data-testid=seci-cond-ceiling]").inner_text()
        assert "INR 4,00,00,000/MW" in page.locator("[data-testid=seci-cond-liquidity]").inner_text()


def test_bid_on_screen(tmp_path):
    with _session() as page:
        _open_tab(page)
        # what you bid
        page.click("[data-testid=seci-step-bid]")
        page.fill("[data-testid=seci-bid-mw]", "200")
        page.click("[data-testid=seci-peak-any]")
        page.click("[data-testid=seci-shortfall-penalty]")
        page.fill("[data-testid=seci-ppa-date]", "2026-06-07")
        dates = page.locator("[data-testid=seci-tab]").inner_text()
        assert "07.06.2028" in dates and "07.12.2027" in dates, "SCSD 24 months after the PPA, Financial Closure 6 months before it"

        # the sources: benchmarks in every empty input, then HiGHS
        page.click("[data-testid=seci-step-sources]")
        page.click("[data-testid=bid-fill]")
        assert page.locator("[data-testid=bid-missing]").count() == 0
        # the losses to the billing meter: each technology's transformer applied after where its CUF is measured
        losses = page.locator("[data-testid=bid-losses]")
        assert "Inverter-duty transformers" in losses.inner_text() and "Power transformer 33/220 kV" in losses.inner_text()
        assert "measured at: Inverter output (LV AC)" in page.locator("[data-testid=source-solar]").inner_text()
        assert "cr/MWp" in page.locator("[data-testid=source-solar]").inner_text(), "capex per MWac, and per MWp from the DC:AC"
        page.click("[data-testid=bid-loss-kind-wind]")
        assert "WTG step-up transformers" in losses.inner_text()
        page.click("[data-testid=bid-loss-kind-bess]")
        assert "MV transformers" in losses.inner_text()
        page.click("[data-testid=bid-size]")
        sizes = page.wait_for_selector("[data-testid=bid-sizes]", timeout=1_800_000)
        page.wait_for_timeout(1500)
        size_text = sizes.inner_text().upper()
        for label in ("SOLAR", "WIND", "BATTERY"):
            assert label in size_text, size_text
        table = page.locator("[data-testid=bid-compliance]").inner_text()
        assert "of hours" in table and "Peak Hours floor" in table, table
        # the hour under the mouse as a see-through Sankey: sources, uses, and the block's floor
        chart = page.locator(".rtc-card", has_text="Year-1 dispatch").locator("canvas").first
        chart.scroll_into_view_if_needed()
        box = chart.bounding_box()
        page.mouse.move(box["x"] + box["width"] * 0.45, box["y"] + box["height"] * 0.5)
        page.wait_for_selector("svg[data-sankey]", timeout=10_000)
        tip = page.locator("svg[data-sankey]").first.locator("xpath=../..").inner_text()
        assert ("Peak block" in tip or "Off-peak block" in tip) and "To the PPA" in tip and "floor" in tip, tip
        # the agent's charts, drawn from the sized plant (the agent opens the same ones as it talks)
        page.click("[data-testid=agent-open]")
        for kind in ("peak_window", "typical_day", "monthly_energy", "irr_vs_tariff", "cash_flows", "money"):
            page.click(f"[data-testid=agent-chart-{kind}]")
            chart = page.locator("[data-testid=agent-chart]")
            assert chart.get_attribute("data-kind") == kind and chart.locator("canvas").count() >= 1, kind
        page.click("[data-testid=agent-chart-peak_window]")
        assert "18:00–10:00" in page.locator("[data-testid=agent-chart]").inner_text()
        page.click("[data-testid=agent-chart-size]")
        assert "big" in page.locator("[data-testid=agent-chart]").get_attribute("class")
        page.click("[data-testid=agent-chart-size]")
        # the CFO's check, typed in the agent: solar year 1 as MW × CUF × 365 × 24, the trace opens beside the conversation
        page.click("[aria-label='Close the chart']")
        page.click("[data-testid=agent-energy-check]")
        page.wait_for_selector("[data-testid=trace-explorer]")
        assert page.locator("[data-testid=trace-name]").inner_text().startswith("Solar at the 33 kV bus · Year 1")
        said = page.locator("[data-testid=agent-log] p.agent").last.inner_text()
        assert "megawatts times" in said and "365 days times 24 hours" in said, said
        assert page.locator("[data-testid=gen-months] tbody tr").count() == 12
        model_mu = page.locator("[data-testid=trace-value]").inner_text().split()[0]
        page.fill("[data-testid=agent-energy-mine]", model_mu)
        page.click("[data-testid=agent-energy-check]")
        page.wait_for_function("document.querySelector('[data-testid=agent-log] p.agent:last-of-type')?.textContent.startsWith('Your')")
        assert "matches the model" in page.locator("[data-testid=agent-log] p.agent").last.inner_text()
        # a second question moves the open model to its year
        page.fill("[data-testid=agent-energy-year]", "5")
        page.click("[data-testid=agent-energy-check]")
        page.wait_for_function("document.querySelector('[data-testid=trace-name]')?.textContent.endsWith('Year 5')")
        assert "year-5 degradation" in page.locator("[data-testid=agent-log] p.agent").last.inner_text()
        page.keyboard.press("Escape")
        page.click("[aria-label='Close the agent']")
        page.click("[data-testid=seci-step-bid]")
        money = page.locator("[data-testid=seci-money]").inner_text()
        assert "EMD" in money.upper() and "cr" in money, money

        # the bid tariff for the target IRR
        page.click("[data-testid=seci-step-finance]")
        page.wait_for_selector("[data-testid=bid-finance]")
        bid_tariff = float(re.search(r"BID TARIFF\s*₹([\d.]+)/kWh", page.locator("[data-testid=bid-finance]").inner_text()).group(1))

        # the tariff quoted: below the target IRR it is warned with the equity and project IRR it earns
        page.click("[data-testid=seci-step-finance]")
        page.wait_for_selector("[data-testid=quote-check]")
        finance_text = page.locator("[data-testid=bid-finance]").inner_text()
        assert "PROJECT IRR" in finance_text.upper(), finance_text
        low = f"{bid_tariff - 0.5:.2f}"
        page.fill("[data-testid=quote-check-input]", low)
        warning = page.locator("[data-testid=quote-check-warning]").inner_text()
        assert "below the" in warning and "equity IRR is" in warning and "project IRR" in warning, warning
        assert "%" in page.locator("[data-testid=quote-check]").inner_text()

        # the particulars, asked once; the quote is the same tariff as on Financials
        page.click("[data-testid=seci-step-forms]")
        assert page.locator("[data-testid=seci-form-tariff]").input_value() == low
        assert "below the" in page.locator("[data-testid=seci-quote]").inner_text()
        for key, value in {"bidder": "Acme Renewables Ltd", "address": "New Delhi", "contact": "A. Kumar", "email": "bids@acme.example",
                           "signatory": "A. Kumar, Director", "letterDate": "2026-05-20", "place": "New Delhi",
                           "substation": "Bhadla-III", "substationScd": "2027-12-31", "netWorthCr": "50000",
                           "experience": "Plant A, 180 MW, 2022-05-01"}.items():
            page.fill(f"[data-testid=seci-form-{key}]", value)
        for key, value in {"bidderType": "company", "connectivity": "final", "bessWith": "solar", "emdForm": "bg", "commonShareholding": "no"}.items():
            page.select_option(f"[data-testid=seci-form-{key}]", value)
        quoted = re.search(r"use ₹([\d.]+)/kwh", page.locator("[data-testid=seci-form-use-tariff]").inner_text(), re.I).group(1)
        assert 0 <= float(quoted) - bid_tariff < 0.0151, (quoted, bid_tariff)
        page.click("[data-testid=seci-form-use-tariff]")
        assert page.locator("[data-testid=seci-form-tariff]").input_value() == quoted
        assert page.locator("[data-testid=seci-quote]").get_attribute("role") is None, "at the target tariff there is no warning"
        checks = page.locator("[data-testid=seci-checks]").inner_text()
        for row in ("capacity", "storage", "substation", "networth", "experience", "emd", "pbg", "dates", "formats"):
            assert page.locator(f"[data-testid=seci-check-{row}]").count() == 1, row
        assert "Not met" not in checks, checks

        # the RfS formats: the ones this bid needs, each blank filled from the bid, the particulars or typed, beside the RfS page
        nav = page.locator("[data-testid=seci-forms-nav]").inner_text().upper()
        assert "ENVELOPE I " in nav and "ENVELOPE II" in nav and "NOT NEEDED FOR THIS BID" in nav, nav
        page.click("[data-testid=seci-formsnav-7\\.1]")
        f71 = page.locator("[data-testid=seci-format-7\\.1]")
        assert "needed for this bid" in f71.locator("[data-testid=seci-format-applies]").inner_text().lower()
        assert page.locator("[data-testid=seci-rfs-frame]").get_attribute("src").endswith("RfS.pdf#page=68&view=FitH")
        assert page.locator("[data-testid=seci-blank-7\\.1-contractedMw]").inner_text() == "200"
        solar = re.match(r"(\d+) MWac \((\d+) MWp DC\)", page.locator("[data-testid=seci-blank-7\\.1-solarMw]").inner_text())
        assert solar and f"{solar.group(1)} MWAC" in size_text.replace(",", "") and f"{solar.group(2)} MWP" in size_text.replace(",", ""), "the installed solar from the sizing, in MWac and MWp"
        assert page.locator("[data-testid=seci-blank-7\\.1-emdAmount]").inner_text().startswith("INR ")
        page.fill("[data-testid=seci-blank-7\\.1-refNo]", "ACME/SECI/2026/01")
        page.locator("[data-testid=seci-format-7\\.1] .seci-blank-ctx").nth(5).click()
        assert "#page=6" in page.locator("[data-testid=seci-rfs-frame]").get_attribute("src"), "a blank's words in the RfS open its page"
        _review(page)
        assert "reviewed" in page.locator("[data-testid=seci-format-reviewed]").inner_text().lower()
        assert "reviewed" in page.locator("[data-testid=seci-formsnav-7\\.1]").inner_text().lower()
        page.fill("[data-testid=seci-blank-7\\.1-refNo]", "ACME/SECI/2026/02")
        assert "changed since review" in page.locator("[data-testid=seci-format-review]").inner_text().lower(), "an edit after the review shows"
        _review(page)
        page.click("[data-testid=seci-formsnav-7\\.10]")
        assert page.locator("[data-testid=seci-format-7\\.10]").count() == 1
        assert page.locator("[data-testid=seci-formsnav-7\\.2]").count() == 1 and "not needed" in page.locator("[data-testid=seci-formsnav-7\\.2]").inner_text().lower()

        # the letter: the RfS wording with the blanks filled, a paragraph changed in place and reverted,
        # Word and PDF of this letter and of every letter the bid needs
        page.click("[data-testid=seci-formsnav-7\\.1]")
        page.click("[data-testid=seci-format-mode-letter]")
        letter = page.locator("[data-testid=seci-letter]")
        text = letter.inner_text()
        for needle in ("Acme Renewables Ltd", "ACME/SECI/2026/02", "COVERING LETTER", "200"):
            assert needle in text, needle
        assert "Choose in the blanks: groupDisclosure" in text, "an alternative not yet chosen asks for the choice"
        page.click("[data-testid=seci-format-mode-blanks]")
        page.select_option("[data-testid=seci-blank-7\\.1-groupDisclosure]", label="No group company has bid (Format 7.8)")
        page.click("[data-testid=seci-format-mode-letter]")
        struck = " ".join(letter.locator("s").all_inner_texts())
        assert "26%" in struck or "shareholder" in struck.lower(), "the alternative not chosen is struck out"
        _review(page)
        para = letter.locator(".letter-para:not(.struck)", has_text="having read, examined").locator(".letter-text").first
        block = para.get_attribute("data-testid").rsplit("-", 1)[1]
        para.click()
        page.keyboard.press("End")
        page.keyboard.type(" We confirm this in writing.")
        page.locator("[data-testid=seci-letter] .letter-heading").first.click()
        assert "We confirm this in writing." in page.locator(f"[data-testid=letter-text-{block}]").inner_text()
        assert "changed from the RfS" in page.locator(f"[data-testid=letter-block-{block}]").inner_text()
        assert "changed since review" in page.locator("[data-testid=seci-format-review]").inner_text().lower(), "an edit to the letter needs a new review"
        page.click(f"[data-testid=letter-revert-{block}]")
        assert "We confirm this in writing." not in page.locator(f"[data-testid=letter-text-{block}]").inner_text()
        page.click(f"[data-testid=letter-text-{block}]")
        page.keyboard.press("End")
        page.keyboard.type(" We confirm this in writing.")
        page.locator("[data-testid=seci-letter] .letter-heading").first.click()
        downloads = {}
        for button in ("seci-letter-word", "seci-letter-pdf", "seci-letters-word", "seci-letters-pdf"):
            with page.expect_download(timeout=120_000) as dl:
                page.click(f"[data-testid={button}]")
            downloads[button] = Path(tmp_path) / f"{button}-{dl.value.suggested_filename}"
            dl.value.save_as(str(downloads[button]))
        with zipfile.ZipFile(downloads["seci-letter-word"]) as z:
            xml = z.read("word/document.xml").decode()
        assert "Acme Renewables Ltd" in xml and "We confirm this in writing." in xml and "<w:strike/>" in xml
        with zipfile.ZipFile(downloads["seci-letters-word"]) as z:
            xml = z.read("word/document.xml").decode()
        assert xml.count('w:type="page"') >= 5 and "FORMAT 7.10" in xml.upper()
        with pymupdf.open(str(downloads["seci-letter-pdf"])) as doc:
            pdf = " ".join(p.get_text(flags=0) for p in doc)  # flags=0 writes ligatures (fi, ffi) as letters
        assert "Acme Renewables Ltd" in pdf and "We confirm this in writing." in " ".join(pdf.split())
        with pymupdf.open(str(downloads["seci-letters-pdf"])) as doc:
            assert doc.page_count >= 6
        page.click("[data-testid=seci-format-mode-blanks]")

        with page.expect_download(timeout=60_000) as dl:
            page.click("[data-testid=seci-forms-download]")
        pack = Path(tmp_path) / dl.value.suggested_filename
        dl.value.save_as(str(pack))
        assert pack.name.endswith(".docx")
        with zipfile.ZipFile(pack) as z:
            xml = z.read("word/document.xml").decode()
        for needle in ("Acme Renewables Ltd", "Bhadla-III", "EMD", "Performance Bank Guarantee", "Amendment-02, p. 3", f"{quoted}",
                       "Format 7.1: Covering Letter", "ACME/SECI/2026/02", "Format 7.10", "RfS pp. 68"):
            assert needle in xml, needle
        assert "Format 7.2:" not in xml, "a single company's pack has no consortium Power of Attorney"

        # the financial model: the Live model workbook matches the app
        page.click("[data-testid=seci-step-finance]")
        page.wait_for_timeout(1500)
        finance = page.locator("[data-testid=bid-finance]").inner_text()
        tariff = re.search(r"BID TARIFF\s*₹([\d.]+)/kWh", finance)
        irr = re.search(r"EQUITY IRR\s*([\d.]+)%", finance)
        assert tariff and irr, finance
        with page.expect_download(timeout=120_000) as dl:
            page.click("[data-testid=bid-download]")
        book = Path(tmp_path) / dl.value.suggested_filename
        dl.value.save_as(str(book))
        _check_live_model(book, tariff=float(tariff.group(1)), irr=float(irr.group(1)) / 100, tmp=tmp_path)

        # every PPA year dispatched in 15-minute blocks with the time-block floors, as one workbook
        page.click("[data-testid=seci-step-size]")
        page.click("[data-testid=bid-d15-run]")
        page.wait_for_selector("[data-testid=bid-d15-xlsx]", timeout=1_800_000)
        assert "of blocks" in page.locator("[data-testid=bid-d15-table]").inner_text()
        with page.expect_download(timeout=120_000) as dl:
            page.click("[data-testid=bid-d15-xlsx]")
        xlsx = Path(tmp_path) / "dispatch.xlsx"
        dl.value.save_as(str(xlsx))
        import openpyxl

        wb = openpyxl.load_workbook(xlsx, read_only=True)
        assert wb.sheetnames[:3] == ["About", "Summary", "Monthly"] and len(wb.sheetnames) == 3 + 25
        assert wb["Year 25"].max_row == 8761, "a row per hour, every year"

        # saved with the bid: the library holds the edit, and the letter reopens with it
        page.click("[data-testid=seci-save]")
        page.fill(".modal input[required]", "Acme SECI letters test")
        page.click(".modal button[type=submit]")
        page.wait_for_selector("[data-testid=seci-saved]")
        page.goto(f"{URL}/scenarios?module=seci")
        page.locator("tr", has_text="Acme SECI letters test").first.locator("a:has-text('Open')").click()
        page.wait_for_selector("[data-testid=seci-tab]", timeout=30_000)
        assert "acme seci letters test" in page.locator("[data-testid=seci-linked]").inner_text().lower()
        page.click("[data-testid=seci-step-forms]")
        page.click("[data-testid=seci-formsnav-7\\.1]")
        page.click("[data-testid=seci-format-mode-letter]")
        assert "We confirm this in writing." in page.locator(f"[data-testid=letter-text-{block}]").inner_text(), "the edit is saved with the bid"


def test_agent_panel():
    with _session(permissions=["microphone"]) as page:
        _open_tab(page)
        info = page.evaluate("fetch('/api/seci/live').then((r) => r.json())")
        assert info["model"] == "gpt-live-1"
        page.click("[data-testid=agent-open]")
        assert "GPT-LIVE" in page.locator("[data-testid=agent-panel]").inner_text().upper()
        if not info["available"]:
            assert "OPENAI_API_KEY" in page.locator("[data-testid=agent-unavailable]").inner_text(), "no key: the panel says how to set it up"
            assert page.locator("[data-testid=agent-start]").is_disabled()
            return
        page.click("[data-testid=agent-start]")
        page.wait_for_selector("[data-testid=agent-error], [data-testid=agent-stop]", timeout=30_000)
        if page.locator("[data-testid=agent-error]").count():
            # a stand-in for OpenAI (or a wrong key) refuses: the panel says so and nothing is left running
            assert "voice service" in page.locator("[data-testid=agent-error]").inner_text()
            assert page.locator("[data-testid=agent-start]").count() == 1
        else:
            page.click("[data-testid=agent-stop]")
