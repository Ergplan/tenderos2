"""Time-block supply floors in the sizing LP (SECI RTC-TM, Amendment-01 cl. 8.1): a floor in every
hour of the peak hours and of the off-peak hours, hard or with a shortfall penalty of
penaltyMultiple x the year's tariff per kWh short."""
from __future__ import annotations

import numpy as np
import pytest

from engine.sizing import dispatch15 as D
from engine.sizing import rtc_lp as M
from tests.engine.test_bid_lp import BESS0, FIN0, HOUR
from tests.engine.test_rtc_lp import COSTS, _case, _formula_delivered

pytestmark = pytest.mark.skipif(M.highspy is None, reason="highspy not installed")

PEAK = ((HOUR >= 18) & (HOUR < 24)).astype(float).tolist()  # six evening hours


def _rules(mult=None):
    extra = {} if mult is None else {"penaltyMultiple": mult}
    return [
        {"id": "peak", "label": "Peak blocks", "target": 0.9, "hours": "peak", "basis": "block", **extra},
        {"id": "offpeak", "label": "Off-peak blocks", "target": 0.7, "hours": "offpeak", "basis": "block", **extra},
    ]


def _payload(rules, **extra):
    solar, wind, _ = _case()
    payload = {
        "ctx": {"demand": [100.0] * M.HOURS, "solarCf": solar.tolist(), "windCf": wind.tolist(), "plantMw": 100},
        "bess": BESS0, "costs": COSTS, "fin": FIN0, "dfrTarget": 0.0, "dfrBasis": "annual",
        "compliance": rules, "peakMask": PEAK, "tariffGuess": 6.0, "returnHourly": True, "returnLifetime": True,
    }
    payload.update(extra)
    return payload


@pytest.fixture(scope="module")
def hard():
    return M.solve(_payload(_rules()))


def test_hard_block_floors_hold_in_every_hour(hard):
    h = hard["hourly"]
    delivered = np.asarray(h["direct"]) + np.asarray(h["discharge"])
    floor = np.where(np.asarray(PEAK) > 0.5, 90.0, 70.0)
    assert (delivered >= floor - 0.02).all(), "every hour at or above its floor (hourly values are rounded to 0.01 MW)"
    for rule in hard["perYear"][0]["rules"]:
        assert rule["basis"] == "block" and rule["met"] and rule["hoursMet"] == 1.0
        assert rule["shortfallMu"] == pytest.approx(0, abs=1e-6)
    assert hard["perYear"][0]["blockShortfallMu"] == pytest.approx(0, abs=1e-6)
    assert all(row["blockShortfallMwh"] == pytest.approx(0, abs=1e-3) for row in hard["lifetime"])


@pytest.mark.slow  # three more solves: run with `make test-slow`
def test_a_penalty_lets_the_plant_fall_short_when_that_is_cheaper(hard):
    free = M.solve(_payload(_rules(0.0)))  # no penalty: shortfall costs nothing
    priced = M.solve(_payload(_rules(1.5)))  # SECI: 1.5 x tariff on the energy short
    steep = M.solve(_payload(_rules(50.0)))  # dearer than any plant: back to the hard floors
    tol = 0.003  # HiGHS gap tolerance
    assert free["tariff"] <= priced["tariff"] + tol <= hard["tariff"] + 2 * tol
    assert free["perYear"][0]["blockShortfallMu"] > priced["perYear"][0]["blockShortfallMu"] - 1e-6
    # a steep penalty leaves only the few hours dearer to cover than ~Rs 300/kWh short
    floor_mu = (0.9 * sum(PEAK) + 0.7 * (M.HOURS - sum(PEAK))) * 100 / 1000
    assert steep["perYear"][0]["blockShortfallMu"] <= min(priced["perYear"][0]["blockShortfallMu"] + 1e-6, 0.001 * floor_mu)
    assert priced["tariff"] - tol <= steep["tariff"] <= hard["tariff"] + tol
    # the reported shortfall is the hourly gap below each hour's floor
    for res in (free, priced):
        h = res["hourly"]
        delivered = np.asarray(h["direct"]) + np.asarray(h["discharge"])
        floor = np.where(np.asarray(PEAK) > 0.5, 90.0, 70.0)
        gap = np.maximum(0.0, floor - delivered).sum() / 1000
        assert res["perYear"][0]["blockShortfallMu"] == pytest.approx(gap, abs=0.05)
        rules = {r["id"]: r for r in res["perYear"][0]["rules"]}
        assert rules["peak"]["shortfallMu"] + rules["offpeak"]["shortfallMu"] == pytest.approx(res["perYear"][0]["blockShortfallMu"], abs=1e-6)
        assert rules["peak"]["penaltyMultiple"] in (0.0, 1.5)


@pytest.mark.parametrize("rules, extra, message", [
    ([{"id": "a", "target": 0.8, "hours": "all", "basis": "annual", "penaltyMultiple": 1.5}], {}, "penaltyMultiple"),
    ([{"id": "a", "target": 0.8, "hours": "all", "basis": "block", "penaltyMultiple": -1}], {}, "penaltyMultiple"),
    ([{"id": "a", "target": 0.8, "hours": "offpeak", "basis": "block"}], {"peakMask": [0.0] * M.HOURS}, "peakMask"),
    ([{"id": "a", "target": 0.8, "hours": "all", "basis": "daily"}], {}, "basis"),
])
def test_block_rule_inputs_are_checked(rules, extra, message):
    with pytest.raises(M.LpInputError, match=message):
        M.solve(_payload(rules, **extra))


def test_fifteen_minute_dispatch_holds_the_block_floors_and_writes_the_workbook(tmp_path):
    """The 25-year dispatch for a time-block tender: each 15-minute block of the peak and off-peak
    hours on its own, a shortfall priced at 1.5 x the tariff, and the workbook adding up."""
    import openpyxl

    from engine.sizing import dispatch15 as D

    payload = _payload(_rules(1.5))
    payload["ctx"] = {**payload["ctx"], "trafoLossPct": 0.004, "lineLossPct": 0.005}
    sized = M.solve(payload)
    assert sized["ok"]
    out = D.run(payload, sized["sizes"], sized["tariff"], years=[1, 2], zip_path=tmp_path / "d.zip", xlsx_path=tmp_path / "d.xlsx")
    for y, lp_year in zip(out["years"], [sized["perYear"][0], None]):
        rules = {r["id"]: r for r in y["rules"]}
        assert set(rules) == {"peak", "offpeak"} and all(r["basis"] == "block" and r["penaltyMultiple"] == 1.5 for r in rules.values())
        # within an hour the 15-minute blocks differ (solar ramps), so a block can miss a floor the
        # hour's average met: the 15-minute check is a little stricter than the hourly sizing
        if lp_year:
            hourly = {r["id"]: r for r in lp_year["rules"]}
            for rid in rules:
                assert rules[rid]["hoursMet"] >= hourly[rid]["hoursMet"] - 0.03, (rid, rules[rid], hourly[rid])
            assert y["shortfallMu"] == pytest.approx(sum(r["shortfallMu"] for r in hourly.values()), rel=0.1)
            assert y["deliveredMu"] == pytest.approx(lp_year["deliveredMu"], rel=0.01)
        assert y["shortfallMu"] == pytest.approx(sum(r["shortfallMu"] for r in rules.values()))
    wb = openpyxl.load_workbook(tmp_path / "d.xlsx", read_only=True)
    assert wb.sheetnames == ["About", "Summary", "Monthly", "Year 01", "Year 02"]
    ws = wb["Year 01"]
    rows = list(ws.iter_rows(values_only=True))
    head, body = rows[0], rows[1:]
    assert len(body) == M.HOURS and body[0][:2] == ("01-Jan", "00:00") and body[-1][:2] == ("31-Dec", "23:00")
    col = {name: i for i, name in enumerate(head)}
    y1 = out["years"][0]
    delivered = sum(r[col["Delivered to PPA (MWh)"]] for r in body) / 1000
    short = sum(r[col["Short of the floor (MWh)"]] for r in body) / 1000
    assert delivered == pytest.approx(y1["deliveredMu"], rel=1e-5)
    assert short == pytest.approx(y1["shortfallMu"], rel=1e-4, abs=1e-6)
    assert "Power transformer and line losses (MWh)" in col and "Time-block floor (MWh)" in col
    summary = list(wb["Summary"].iter_rows(values_only=True))
    assert [r[0] for r in summary[1:]] == [1, 2]
    monthly = list(wb["Monthly"].iter_rows(values_only=True))
    assert len(monthly) == 1 + 24


SOLAR = ((HOUR >= 8) & (HOUR < 18)).astype(float).tolist()  # Solar Hours 08:00-18:00


def _amd3(share=0.5):
    """Amendment-03 (SECI FDRE-RTC-V): 90% in the peak blocks, 50-60% in the Solar Hours (power above
    60% earns at most half the tariff), 70% in the Remaining Hours, shortfall at 1.5 x tariff."""
    return [
        {"id": "peak", "label": "Peak", "target": 0.9, "hours": "peak", "basis": "block", "penaltyMultiple": 1.5},
        {"id": "solar", "label": "Solar Hours", "target": 0.5, "max": 0.6, "excessTariffShare": share, "hours": "solar", "basis": "block", "penaltyMultiple": 1.5},
        {"id": "remaining", "label": "Remaining Hours", "target": 0.7, "hours": "remaining", "basis": "block", "penaltyMultiple": 1.5},
    ]


@pytest.mark.slow  # two solves and a 15-minute dispatch
def test_amendment_03_caps_the_solar_hours_and_prices_their_excess():
    fin = {**FIN0, "sellSurplus": True, "surplusPrice": 4.0}  # a market that pays more than half the tariff
    payload = _payload(_amd3(), solarMask=SOLAR, fin=fin, ppaFirst=True, ctx={**_payload([])["ctx"], "sellSurplus": True}, returnDispatch=True)
    res = M.solve(payload)
    assert res["ok"], res.get("status")
    h = res["hourly"]
    delivered = np.asarray(h["direct"]) + np.asarray(h["discharge"])
    solar = (np.asarray(SOLAR) > 0.5) & (np.asarray(PEAK) < 0.5)
    assert delivered[solar].max() <= 60.0 + 0.02, "the PPA takes at most 60% in the Solar Hours"
    assert delivered[~solar].max() > 60.0, "and up to the contract outside them"
    rules = {r["id"]: r for r in res["perYear"][0]["rules"]}
    assert set(rules) == {"peak", "solar", "remaining"}
    # the Solar Hours surplus goes to the buyer at half the tariff (the market pays more), none to the market
    assert res["buyerExcessShare"] == 0.5
    life = res["lifetime"][0]
    assert life["buyerExcessMwh"] > 1000 and life["exportMwh"] == pytest.approx(0, abs=1.0)
    # year 1's hours say what a sale earns at most: half the tariff in a Solar Hour, no cap elsewhere
    assert h["saleCap"][12] == pytest.approx(0.5 * res["tariff"], rel=1e-3) and h["saleCap"][2] is None
    # the workbook's dispatch sheets: a Solar Hours column, the 60% cap and the 50% floor reproduce
    # the LP's energy and shortfall year by year
    d = res["dispatch"]
    assert d["blocks"] and (d["solarFloor"], d["solarMax"], d["offFloor"]) == (0.5, 0.6, 0.7)
    assert sum(d["solarHours"]) == 365 * 10, "08:00-18:00, none of them a Peak Hour (18:00-24:00)"
    lives = {y["year"]: y for y in res["lifetime"]}
    for y in d["byYear"]:
        delivered, short, low = _formula_delivered(d, y, res["sizes"], fin, BESS0)
        assert delivered == pytest.approx(lives[y["year"]]["deliveredMwh"], rel=1e-5)
        assert short == pytest.approx(lives[y["year"]]["blockShortfallMwh"], rel=1e-4, abs=1.0)
        assert low >= -1e-3
    # the 15-minute dispatch keeps the cap and gives the buyer the Solar Hours excess
    y1 = D.run(payload, res["sizes"], res["tariff"], years=[1])["years"][0]
    assert {r["id"] for r in y1["rules"]} == {"peak", "solar", "remaining"}
    assert y1["toBuyerMu"] > 1 and y1["toBuyerMu"] <= y1["soldMu"] + 1e-9
    # with a market below half the tariff the same hours sell to the market
    cheap = M.solve(_payload(_amd3(), solarMask=SOLAR, fin={**fin, "surplusPrice": 1.0}, ppaFirst=True, ctx={**_payload([])["ctx"], "sellSurplus": True}))
    assert cheap["lifetime"][0]["buyerExcessMwh"] == pytest.approx(0, abs=1.0) and cheap["lifetime"][0]["exportMwh"] > 0


def test_amendment_03_rules_are_checked():
    with pytest.raises(M.LpInputError, match="solarMask"):
        M.solve(_payload(_amd3()))
    bad = _amd3()
    bad[1]["max"] = 0.4  # below its own floor
    with pytest.raises(M.LpInputError, match="max"):
        M.solve(_payload(bad, solarMask=SOLAR))
